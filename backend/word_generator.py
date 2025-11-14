"""
Word Document Generator
Converts filled HTML template to Word document (.docx)
"""
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from bs4 import BeautifulSoup
import re
from typing import Optional
import io


def html_to_word(html_content: str, output_path: Optional[str] = None) -> bytes:
    """
    Convert HTML content to Word document
    
    Args:
        html_content: HTML string with filled placeholders
        output_path: Optional path to save file (if None, returns bytes)
        
    Returns:
        bytes: Word document as bytes (if output_path is None)
    """
    # Parse HTML
    soup = BeautifulSoup(html_content, 'html.parser')
    
    # Create Word document
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    
    # Extract body content
    body = soup.find('body') or soup
    
    # Process content recursively
    _process_element(body, doc)
    
    # Save or return bytes
    if output_path:
        doc.save(output_path)
        return None
    else:
        # Save to bytes buffer
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.read()


def _process_element(element, doc):
    """Recursively process HTML elements and add to Word document"""
    
    # Handle text nodes
    if isinstance(element, str):
        if element.strip():
            # Add text to last paragraph or create new one
            if doc.paragraphs:
                doc.paragraphs[-1].add_run(element)
            else:
                p = doc.add_paragraph(element)
        return
    
    # Handle different HTML tags
    tag_name = element.name.lower() if hasattr(element, 'name') else None
    
    if tag_name == 'table':
        _process_table(element, doc)
    elif tag_name == 'p':
        text = element.get_text(strip=True)
        if text:
            p = doc.add_paragraph(text)
            # Handle alignment
            if element.get('align'):
                align_map = {
                    'center': WD_ALIGN_PARAGRAPH.CENTER,
                    'right': WD_ALIGN_PARAGRAPH.RIGHT,
                    'left': WD_ALIGN_PARAGRAPH.LEFT
                }
                p.alignment = align_map.get(element.get('align').lower(), WD_ALIGN_PARAGRAPH.LEFT)
    elif tag_name in ['h1', 'h2', 'h3', 'h4', 'h5', 'h6']:
        text = element.get_text(strip=True)
        if text:
            heading_level = int(tag_name[1]) if len(tag_name) > 1 else 1
            doc.add_heading(text, level=min(heading_level, 9))
    elif tag_name == 'br':
        if doc.paragraphs:
            doc.paragraphs[-1].add_run().add_break()
        else:
            doc.add_paragraph()
    elif tag_name in ['div', 'span', 'body']:
        # Process children
        for child in element.children:
            _process_element(child, doc)
    elif tag_name in ['ul', 'ol']:
        _process_list(element, doc)
    elif tag_name == 'li':
        text = element.get_text(strip=True)
        if text:
            p = doc.add_paragraph(text, style='List Bullet')
    else:
        # For other tags, just get text content
        text = element.get_text(strip=True)
        if text and tag_name not in ['script', 'style', 'head']:
            if doc.paragraphs:
                doc.paragraphs[-1].add_run(text)
            else:
                doc.add_paragraph(text)


def _process_table(table_element, doc):
    """Process HTML table and add to Word document"""
    rows = table_element.find_all('tr')
    if not rows:
        return
    
    # Determine number of columns
    max_cols = 0
    for row in rows:
        cells = row.find_all(['td', 'th'])
        max_cols = max(max_cols, len(cells))
    
    if max_cols == 0:
        return
    
    # Create Word table
    word_table = doc.add_table(rows=len(rows), cols=max_cols)
    word_table.style = 'Table Grid'
    
    # Fill table cells
    for row_idx, row in enumerate(rows):
        cells = row.find_all(['td', 'th'])
        for col_idx, cell in enumerate(cells):
            if col_idx < max_cols:
                cell_text = cell.get_text(strip=True)
                word_table.rows[row_idx].cells[col_idx].text = cell_text
                
                # Handle cell alignment
                if cell.get('align'):
                    align_map = {
                        'center': WD_ALIGN_PARAGRAPH.CENTER,
                        'right': WD_ALIGN_PARAGRAPH.RIGHT,
                        'left': WD_ALIGN_PARAGRAPH.LEFT
                    }
                    para = word_table.rows[row_idx].cells[col_idx].paragraphs[0]
                    para.alignment = align_map.get(cell.get('align').lower(), WD_ALIGN_PARAGRAPH.LEFT)


def _process_list(list_element, doc):
    """Process HTML list (ul/ol) and add to Word document"""
    items = list_element.find_all('li', recursive=False)
    for item in items:
        text = item.get_text(strip=True)
        if text:
            # Determine if ordered or unordered
            if list_element.name == 'ol':
                # For ordered lists, we'd need to track numbering
                # For now, use bullet style
                p = doc.add_paragraph(text, style='List Bullet')
            else:
                p = doc.add_paragraph(text, style='List Bullet')

